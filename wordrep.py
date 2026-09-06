"""Lee un informe Word subido a Drive: extrae un resumen (HTML) y devuelve el
archivo en base64 para incrustarlo cifrado y permitir su descarga tras login."""
import re, base64, html as _html
def extract(path):
    import docx
    d = docx.Document(path)
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    summ = []; started = False
    for t in paras:
        up = t.upper()
        if not started:
            if "RESUMEN" in up: started = True
            continue
        if re.match(r'^\d+\.\s', t) or re.match(r'^\d+\.\d', t):
            if summ: break
            continue
        summ.append(t)
        if sum(len(x) for x in summ) > 1500: break
    # Los titulos de tabla ("Indicadores clave - 8.o corte vs. 7.o") viven en
    # d.paragraphs, pero la tabla que los sigue NO: si el informe cierra el
    # resumen con ellos, quedan como encabezados huerfanos sin nada debajo.
    # Se recortan los del final; los que van entre parrafos de texto se quedan,
    # porque ahi si rotulan lo que viene despues.
    _titulo = lambda t: len(t) < 90 and t[-1] not in ".!?:;"
    while summ and _titulo(summ[-1]): summ.pop()
    if not summ:
        summ = [p for p in paras if len(p) > 40][:5]
    body = "".join(f"<p>{_html.escape(x)}</p>" for x in summ)
    b64 = base64.b64encode(open(path, "rb").read()).decode()
    return body, b64
